"""Exportação do relatório: XLSX, CSV e PDF."""
from __future__ import annotations

import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

from .report_generator import Report
from ..storage.atomic import atomic_output

def _safe_csv_cell(value):
    """Neutralize spreadsheet formula injection in untrusted textual report fields."""
    if not isinstance(value, str):
        return value
    trimmed = value.lstrip(" \t\r\n")
    if value.startswith(("\t", "\r", "\n")) or (trimmed and trimmed[0] in "=+-@"):
        return "'" + value
    return value

HEADERS = ["Cidade", "Empresa", "Competência", "Status", "XML encontrados", "XML baixados",
           "XML já existentes", "XML falhas", "Livro Fiscal", "Mensagem"]


def _rows(rep: Report) -> list[list]:
    return [[r.cidade, r.cnpj_usuario, r.competencia, r.status.value, r.xml_found, r.xml_downloaded,
             r.xml_existing, r.xml_failed, "OK" if r.livro_ok else "—", r.message or ""] for r in rep.runs]


def to_csv(rep: Report, path: Path) -> Path:
    with path.open("w", newline="", encoding="utf-8-sig") as fh:
        w = csv.writer(fh, delimiter=";")
        w.writerow(HEADERS)
        w.writerows([[_safe_csv_cell(value) for value in row] for row in _rows(rep)])
    return path


def to_xlsx(rep: Report, path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Empresas"
    ws.append(HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in _rows(rep):
        ws.append(row)
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = min(60, max(12, max(len(str(c.value or "")) for c in col) + 2))
    summary = wb.create_sheet("Resumo")
    for line in rep.to_text().splitlines():
        summary.append([line])
    cities = wb.create_sheet("Cidades")
    cities.append(["Cidade", "Empresas", "Concluídas", "Pendências", "Erros"])
    for c in rep.cities:
        cities.append([c.city, c.empresas, c.concluidas, c.pendencias, c.erros])
    # Explicit text cells preserve the source text and prevent Excel from executing it.
    for sheet in wb:
        for cells in sheet:
            for cell in cells:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
    wb.save(path)
    return path


def to_pdf(rep: Report, path: Path) -> Path:
    styles = getSampleStyleSheet()
    doc = SimpleDocTemplate(str(path), pagesize=A4, title="NexFiscal — Relatório final")
    story = [Paragraph("NexFiscal — Relatório final", styles["Title"]), Spacer(1, 8)]
    for line in rep.to_text().splitlines():
        if line.strip("= "):
            story.append(Paragraph(line.replace("&", "&amp;").replace("<", "&lt;"), styles["Normal"]))
        else:
            story.append(Spacer(1, 6))
    doc.build(story)
    return path


def export(rep: Report, path: Path) -> Path:
    ext = path.suffix.lower()
    writers = {".csv": to_csv, ".xlsx": to_xlsx, ".pdf": to_pdf}
    if ext not in writers:
        raise ValueError("Formato não suportado (use .xlsx, .csv ou .pdf).")
    with atomic_output(path) as temporary:
        writers[ext](rep, temporary)
    return path
